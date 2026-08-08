"""Polish Phase-3 DOCX: GitHub URL + Algorithm 1 line breaks."""
from __future__ import annotations

from pathlib import Path

from docx import Document

PATH = Path(r"D:\_August2026\KanakWsnPaper\Reference-FreeRelativeLocalizationWsn.docx")
MIRROR = Path(
    r"D:\_August2026\KanakWsnPaper\kanakpatel-main\WsnLocalization2026\paper"
    r"\Direction-Aware Routing via Reference-Free Relative Localization in WSN V04.docx"
)
GITHUB = "https://github.com/hsmazumdar/kanakpatel/tree/main/WsnLocalization2026"

ALGO_LINES = [
    "Algorithm 1 Asymmetric neighbor attraction with periodic normalization",
    "Input: communication graph G=(V,E); expected distances d*_ij for stored neighbor pairs; "
    "attraction beta in (0,1); neighbor cap k; radio range R (Tx); moves_per_tick; T_max; stop rule",
    "Output: virtual positions {p_i} up to similarity",
    "1: Initialize p_i for all i in V",
    "   (density-Tx study: random or scrambled; topology-quality study: classical MDS on "
    "shortest-path completion of expected distances)",
    "2: for t <- 1 to T_max do",
    "3:     for m <- 1 to moves_per_tick do",
    "4:         Sample a node u in V",
    "5:         for each stored neighbor v of u do",
    "6:             if ||p_u - p_v|| > d*_uv then",
    "7:                 p_v <- p_v + beta (p_u - p_v)   // move only v (asymmetric update)",
    "8:     Optionally Normalize({p_i}) to a fixed canvas span   // simulation convenience",
    "9:     if hard or soft stopping criterion is satisfied then break",
    "10: return {p_i}",
]


def set_text(para, text: str) -> None:
    if para.runs:
        para.runs[0].text = text
        for r in para.runs[1:]:
            r.text = ""
    else:
        para.add_run(text)


def set_multiline(para, lines: list[str]) -> None:
    set_text(para, "")
    run = para.runs[0] if para.runs else para.add_run()
    run.text = lines[0]
    for line in lines[1:]:
        run.add_break()
        run.add_text(line)


def main() -> None:
    doc = Document(str(PATH))
    notes = []

    for p in doc.paragraphs:
        t = p.text.strip()
        if "github.com/hsmazumdar/kanakpatel" in t and (
            t.startswith("http") or "http" in t and len(t) < 200
        ):
            # Corrupt concat or short URL
            if t != GITHUB:
                set_text(p, GITHUB)
                notes.append("fixed GitHub URL paragraph")
        if t.startswith("Algorithm 1 Asymmetric neighbor attraction"):
            set_multiline(p, ALGO_LINES)
            notes.append("reformatted Algorithm 1 with line breaks")

    # Table 6 finding cell
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "Directional consistency under planned-geometry validation" in p.text:
                        set_text(p, "Directional consistency under planned geometry")
                        notes.append("normalized Table 6 finding label")
                    elif "Directional consistency under Mode A" in p.text:
                        set_text(p, "Directional consistency under planned geometry")
                        notes.append("fixed Table 6 Mode A label")

    doc.save(str(PATH))
    doc.save(str(MIRROR))

    # Verify
    doc2 = Document(str(PATH))
    text = "\n".join(p.text for p in doc2.paragraphs)
    for table in doc2.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    checks = {
        "Mode A": text.count("Mode A"),
        "Mode B": text.count("Mode B"),
        "NS-3": text.count("NS-3"),
        "bad_url_concat": text.count("kanakpatelhttps"),
        "good_url": text.count(GITHUB),
        "Algorithm 1": text.count("Algorithm 1"),
        "[21]": text.count("[21]"),
        "[35]": text.count("[35]"),
        "Relevance to IoT": text.count("Relevance to IoT"),
    }
    report = PATH.parent / "PHASE3_MANUSCRIPT_EDIT.txt"
    prev = report.read_text(encoding="utf-8") if report.exists() else ""
    extra = (
        "\nPOLISH PASS\n-----------\n"
        + "\n".join(f"- {n}" for n in notes)
        + "\nVerify: "
        + ", ".join(f"{k}={v}" for k, v in checks.items())
        + "\n"
    )
    report.write_text(prev.rstrip() + "\n" + extra, encoding="utf-8")
    print("OK", checks)
    for n in notes:
        print("-", n)


if __name__ == "__main__":
    main()
