"""
Generate revised manuscript V02 as a Word document.
Removes unfulfilled / inconsistent claims; adds density–Tx experimental findings.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(
    r"D:\_July2026\KanakPaperLocalization"
    r"\Direction-Aware Routing via Reference-Free Relative Localization in WSN V02.docx"
)
FIGS = Path(r"D:\_July2026\KanakPaperLocalization\Results\figures")


def set_run_font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text, *, style=None, size=11, bold=False, italic=False, space_after=6, align=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_heading_custom(doc, text, level=1):
    # Use built-in heading styles for outline, then restyle font
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=14 if level == 1 else 12, bold=True)
    return p


def add_caption(doc, text):
    p = add_para(doc, text, size=10, italic=True, space_after=10)
    return p


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=9, bold=True)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = ""
            p = cells[c_i].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=9)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def try_add_figure(doc, path: Path, width=5.8):
    if path.exists():
        doc.add_picture(str(path), width=Inches(width))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return True
    return False


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # ----- Title -----
    add_para(
        doc,
        "Direction-Aware Routing via Reference-Free Relative Localization in Wireless Sensor Networks",
        size=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=4,
    )
    add_para(
        doc,
        "Version 02 — Revised manuscript (claims reconciled; density–Tx deployment study added)",
        size=11,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=12,
    )
    add_para(
        doc,
        "Kanak Patel1* and Himanshu S. Mazumdar2",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=2,
    )
    add_para(
        doc,
        "1,2 Research and Development Center, Dharmsinh Desai University, Nadiad 387001, Gujarat, India",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=2,
    )
    add_para(
        doc,
        "*Corresponding author: kanakpatel.rnd@ddu.ac.in  |  Contributing: hsmazumdar@ddu.ac.in",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=14,
    )

    # ----- Abstract -----
    add_heading_custom(doc, "Abstract", 1)
    add_para(
        doc,
        "This paper revisits a practical question for wireless sensor networks (WSNs): is full "
        "metric localization necessary for direction-aware routing? We argue that directional "
        "consistency—not absolute position accuracy—is the functional requirement for greedy "
        "forwarding. We present a lightweight, reference-free relative localization procedure "
        "based on expected distance constraints and an asymmetric neighbor-attraction update. "
        "Local updates use only neighbor state; however, the present simulator applies a global "
        "scale-normalization step for numerical stability, which we explicitly treat as a "
        "simulation convenience rather than a claim of fully decentralized operation. "
        "Convergence is established locally (each selected violated edge strictly shrinks) and "
        "treated globally as an empirically observed relaxation process, not a completed "
        "Lyapunov proof. Beyond prior directional-consistency experiments, this revision adds a "
        "controlled density-versus-transmission-range study (600 trials) showing that higher "
        "node density enables substantially lower radio range—and thus lower transmit energy—"
        "while maintaining a usable neighbor degree (≈8–10). On an 800×800 canvas with neighbor "
        "cap k=10, the minimum Tx (% of diagonal) meeting degree ≥8 and full connectivity falls "
        "from 20% at N=50 to 8% at N=300, reducing a path-loss energy proxy (η=2.5) by about "
        "an order of magnitude. The results support energy-aware deployment guidance: raise "
        "density to reduce Tx, rather than over-provisioning radio range.",
        size=11,
        space_after=8,
    )
    add_para(
        doc,
        "Keywords: Wireless sensor networks; relative localization; direction-aware routing; "
        "node density; transmission range; energy-aware deployment; anchor-free embedding.",
        size=10,
        italic=True,
        space_after=12,
    )

    # ----- Revision note -----
    add_heading_custom(doc, "Revision note (V01 → V02)", 1)
    add_para(
        doc,
        "This Version 02 responds to editorial and reviewer concerns by (i) removing or "
        "softening claims that were not supported by a consistent, auditable experiment "
        "matrix; (ii) reframing convergence as local lemma + empirical global behavior; "
        "(iii) acknowledging global normalization as a simulation limitation; (iv) removing "
        "unsupported scalability extrapolations (e.g., 10,000-node projections and "
        "inconsistent iteration-count statements such as ~4000 iterations at N=1000); "
        "and (v) adding a new, fully logged density-versus-Tx experimental section with CSV "
        "artifacts under Results/. Dummy “implemented gossip normalization” and packet-level "
        "routing claims that were not executed in the current codebase are not asserted here; "
        "they remain future work.",
        size=10,
        space_after=12,
    )

    # ----- 1 Introduction -----
    add_heading_custom(doc, "1 Introduction", 1)
    add_para(
        doc,
        "In WSNs, geographic or direction-aware routing (e.g., GPSR [1]) has long relied on "
        "node positions, so localization is often treated as a prerequisite. Prior work shows "
        "greedy routing can tolerate position error if directional relationships remain "
        "consistent [14,15]. Nonetheless, many localization schemes still optimize metric "
        "accuracy and introduce anchors, ranging hardware, or centralized computation [2–7].",
    )
    add_para(
        doc,
        "We study a lighter objective: construct a relative embedding that preserves topology "
        "and direction up to similarity (translation, rotation, uniform scale), using only "
        "expected distance constraints and local asymmetric updates. The intended consumer is "
        "direction-aware forwarding, not surveying-grade localization.",
    )
    add_para(doc, "Contributions of this revised manuscript:", bold=True)
    bullets = [
        "Routing-centric formulation — directional consistency as the primary objective for greedy forwarding.",
        "Asymmetric neighbor-attraction algorithm with a rigorous local reduction lemma and an explicitly empirical global convergence statement.",
        "Honest distributed-scope statement — local updates are neighbor-only; global normalization in the current simulator is disclosed as a limitation.",
        "New density–Tx deployment study — quantified trade-off showing denser deployments permit lower Tx (energy) while preserving usable degree and connectivity.",
        "Reproducible experiment artifacts — parameter matrix, raw/summary CSVs, figures, and captions under Results/.",
    ]
    for b in bullets:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(b)
        set_run_font(run, size=11)

    add_para(
        doc,
        "Organization: Section 2 reviews related work. Section 3 states the model. Section 4 "
        "describes the algorithm. Section 5 presents convergence status. Section 6 reports "
        "experiments, including the new density–Tx study. Section 7 discusses routing use. "
        "Section 8 concludes with limitations and future work.",
        space_after=10,
    )

    # ----- 2 Related -----
    add_heading_custom(doc, "2 Related Work", 1)
    add_para(
        doc,
        "Range-based methods [2–4] estimate distances from RSS/ToA/TDoA but are sensitive to "
        "multipath and typically need anchors or calibrated hardware. Range-free methods [5–7] "
        "reduce hardware cost yet often still depend on anchors and idealized density. Relative "
        "and MDS-style embeddings [4,8–10] remove absolute anchors but frequently retain "
        "distance estimates or centralized computation. Virtual coordinate systems [16–18] "
        "target routing stretch and delivery with abstract coordinates. Our focus is closer to "
        "routing-oriented relative maps: enforce soft distance-order constraints sufficient for "
        "angular consistency, without claiming metric reconstruction.",
    )
    add_para(
        doc,
        "Table 1 compares paradigms along telecom-relevant axes. Unlike V01, we do not claim "
        "that the proposed method is “fully distributed” in the strict sense while a global "
        "normalization step is used in simulation.",
    )
    add_caption(
        doc,
        "Table 1. Comparison of localization approaches along telecommunication-relevant axes "
        "(revised: Proposed Method marked Partial for distribution due to simulated global normalization).",
    )
    add_table(
        doc,
        [
            "Approach",
            "Anchors?",
            "Ranging?",
            "Distributed?",
            "Training?",
            "Routing-oriented?",
        ],
        [
            ["Range-based [2–4]", "Yes", "Yes", "Partial", "No", "No"],
            ["Range-free / DV-Hop [5–7]", "Yes", "No", "Yes", "No", "No"],
            ["MDS-based [4,8–10]", "No", "Yes", "Often centralized", "No", "No"],
            ["Learning / GNN-style [11]", "No", "Often yes", "Varies", "Yes", "No"],
            ["Virtual coordinates [16–18]", "No", "No", "Yes", "No", "Partial"],
            ["Proposed (this work)", "No", "No*", "Partial†", "No", "Yes"],
        ],
    )
    add_para(
        doc,
        "*Expected distances in this work come from planned spacing (structured deployments), "
        "not physical ranging hardware. "
        "†Local asymmetric updates are neighbor-only; the current simulator applies global "
        "scale normalization after batches of updates.",
        size=9,
        italic=True,
    )

    # ----- 3 Model -----
    add_heading_custom(doc, "3 System Model and Problem Formulation", 1)
    add_heading_custom(doc, "3.1 Network model", 2)
    add_para(
        doc,
        "We model the network as an undirected communication graph G=(V,E). Each node i∈V "
        "maintains a virtual position p_i∈R². Positions need not equal geographic coordinates. "
        "Nodes exchange state only with graph neighbors. Connectivity of G is assumed for the "
        "analysis of a single connected component; sparse regimes are studied empirically in "
        "Section 6.3.",
    )
    add_heading_custom(doc, "3.2 Expected distance constraints", 2)
    add_para(
        doc,
        "For each edge (or stored neighbor pair) the algorithm uses an expected distance d*_ij. "
        "In this manuscript, expected distances come from planned geometry in structured "
        "deployments (validation baseline). Constraints are taken from reference geometry, "
        "as in the present WinForms/Python simulators that build a k-nearest unit-disk "
        "distance matrix from a grid-with-jitter layout. All reported experimental results "
        "use this planned-geometry constraint source.",
        space_after=8,
    )
    add_heading_custom(doc, "3.3 Problem statement", 2)
    add_para(
        doc,
        "Given G and constraints {d*_ij}, design a reference-free procedure such that virtual "
        "positions become mutually consistent with the constraints up to similarity, using "
        "local interactions, and yielding directional cues adequate for greedy forwarding.",
    )

    # ----- 4 Algorithm -----
    add_heading_custom(doc, "4 Proposed Localization Algorithm", 1)
    add_heading_custom(doc, "4.1 Asymmetric neighbor attraction", 2)
    add_para(
        doc,
        "When node u is selected and neighbor v violates ||p_u−p_v|| > d*_uv, only v is moved "
        "toward u by a fraction β∈(0,1) of the displacement (asymmetric update). Satisfied "
        "constraints are left unchanged. This one-sided correction is simple and matches the "
        "implementation used in our experiments (β=0.3 in the density–Tx study).",
    )
    add_heading_custom(doc, "4.2 Scale normalization (simulation)", 2)
    add_para(
        doc,
        "Because embeddings are defined up to similarity, the simulator periodically rescales "
        "the map to occupy a fixed fraction of the canvas (90% span normalization). This uses "
        "global bounding-box information. We do not claim a deployed gossip implementation in "
        "V02; replacing normalization by local consensus remains future work [19].",
    )
    add_heading_custom(doc, "4.3 Algorithm sketch", 2)
    add_para(
        doc,
        "Initialize positions randomly (or from a shuffled layout). Repeatedly: sample a node; "
        "relax violated neighbor constraints asymmetrically; optionally normalize; stop when "
        "no sampled batch finds a violation (hard stop) or when residual violation falls below "
        "a tolerance (soft stop). Neighbor lists are truncated to the k nearest contacts inside "
        "radio range R (Tx).",
    )

    # ----- 5 Convergence -----
    add_heading_custom(doc, "5 Convergence Status", 1)
    add_para(
        doc,
        "We separate what is proved from what is observed, addressing reviewer concerns that "
        "V01 over-promised Lyapunov-style global guarantees.",
    )
    add_heading_custom(doc, "5.1 Lemma 1 (local constraint reduction)", 2)
    add_para(
        doc,
        "If ||p_u−p_v|| > d*_uv and p_v is updated toward p_u by factor β∈(0,1), then the "
        "Euclidean length of edge (u,v) strictly decreases. Proof sketch: the new distance "
        "equals (1−β)·||p_u−p_v|| < ||p_u−p_v||. Thus every selected violated edge makes "
        "local progress.",
    )
    add_heading_custom(doc, "5.2 Global energy is not claimed monotone", 2)
    add_para(
        doc,
        "Define E(P)=Σ_{(i,j)} max(0, ||p_i−p_j||−d*_ij)^2. A local fix of (u,v) may increase "
        "violations on other edges incident to v. Therefore we do not claim that E decreases "
        "after every elementary update. Empirically, randomized relaxation drives residual "
        "violation downward under consistent constraints, but a full asynchronous stochastic "
        "convergence proof is left open.",
    )
    add_heading_custom(doc, "5.3 Statement of theoretical status", 2)
    add_para(doc, "Rigorous: Lemma 1 (local reduction); E(P)≥0 with equality only if all soft constraints are met.", space_after=4)
    add_para(doc, "Empirical: repeated updates reduce aggregate violation in the logged experiments; hard batch-wise finish flags can be rare in dense, high-degree regimes even when residuals are small—hence soft criteria in Section 6.4.", space_after=4)
    add_para(doc, "Not claimed in V02: global Lyapunov decrease; uniqueness of the global minimizer; fully decentralized normalization.", space_after=10)

    # ----- 6 Experiments -----
    add_heading_custom(doc, "6 Experimental Evaluation", 1)
    add_heading_custom(doc, "6.1 Simulation setup (auditable matrix)", 2)
    add_para(
        doc,
        "The density–Tx study uses a headless Python controller mirroring the WsnMap "
        "distance-constraint mode. Fixed controls: canvas 800×800 px; attraction β=0.3; "
        "100 random node updates per tick; maximum 250 ticks; path-loss exponent η=2.5 for "
        "the energy proxy (R/diag)^η. Swept factors:",
    )
    add_table(
        doc,
        ["Factor", "Values", "Role"],
        [
            ["N (nodes)", "50, 100, 150, 200, 300", "Density on fixed area"],
            ["Tx (% diagonal)", "5, 8, 10, 12, 15, 20, 25, 30", "Radio range / Tx proxy"],
            ["k (neighbor cap)", "6, 10, 12", "Stored neighbors (txRange)"],
            ["Seeds", "0…4", "Five Monte-Carlo repeats"],
        ],
    )
    add_caption(
        doc,
        "Table 2. Input parameter matrix for the density–versus–transmission-range study "
        "(600 trials). Artifacts: Results/param_matrix.csv and Results/results_raw.csv.",
    )
    add_para(
        doc,
        "Nodes are placed on an approximately square grid with intra-cell jitter. Radio "
        "neighbors are all nodes within Tx distance, truncated to the k nearest. Soft "
        "success requires: connected graph (no isolates), mean stored degree ≥6, and mean "
        "constraint violation ≤1 px after the run.",
    )

    add_heading_custom(doc, "6.2 Density versus transmission range (new findings)", 2)
    add_para(
        doc,
        "Lower Tx saves energy only if density keeps average degree high enough for "
        "localization and routing. We therefore minimize Tx subject to mean degree ≥8 and "
        "full connectivity (k=10 primary).",
    )

    try_add_figure(doc, FIGS / "fig1_min_tx_vs_N.png")
    add_caption(
        doc,
        "Fig. 1. Minimum Tx (% of canvas diagonal) versus node count N for mean stored degree "
        "≥8 and a connected graph (k=10). Higher density reduces the feasible Tx percentage.",
    )
    try_add_figure(doc, FIGS / "fig2_energy_vs_N.png")
    add_caption(
        doc,
        "Fig. 2. Transmission-energy proxy (R/diag)^η (η=2.5) at the energy-optimal Tx of "
        "Fig. 1 versus N (k=10). Density increase yields roughly an order-of-magnitude drop "
        "in the proxy from N=50 to N=300.",
    )

    add_caption(
        doc,
        "Table 3. Energy-optimal Tx versus node density (k=10, target mean degree ≥8, "
        "connected). Soft converge rate uses the Section 6.1 criterion.",
    )
    add_table(
        doc,
        ["N", "Density (1/px²)", "Min Tx %", "Energy proxy", "Soft conv.", "Mean degree"],
        [
            ["50", "7.81e-05", "20", "0.0179", "100%", "8.76"],
            ["100", "1.56e-04", "15", "0.0087", "60%", "9.33"],
            ["150", "2.34e-04", "10", "0.0032", "20%", "8.12"],
            ["200", "3.13e-04", "10", "0.0032", "40%", "9.52"],
            ["300", "4.69e-04", "8", "0.0018", "20%", "9.52"],
        ],
    )

    try_add_figure(doc, FIGS / "fig3_degree_vs_tx.png")
    add_caption(
        doc,
        "Fig. 3. Mean stored degree versus Tx for several N (k=10). The dashed reference at "
        "degree 8 is the deployment design target; denser nets reach it at lower Tx.",
    )
    try_add_figure(doc, FIGS / "fig5_connected_vs_tx.png")
    add_caption(
        doc,
        "Fig. 4. Connectedness rate versus Tx for several N (k=10). Connectivity is a "
        "prerequisite for reliable relative localization.",
    )
    try_add_figure(doc, FIGS / "fig4_soft_converge_heatmap.png")
    add_caption(
        doc,
        "Fig. 5. Soft-convergence rate versus N (rows) and Tx (columns), k=10. Soft success "
        "requires connectivity, degree ≥6, and residual violation ≤1 px.",
    )

    add_para(
        doc,
        "Deployment takeaway: raising density allows cutting radio range while holding degree "
        "near 8–10. Over-provisioning Tx in dense networks wastes energy and admits weak, "
        "long links that add little localization value once k neighbors are already available.",
        space_after=8,
    )

    add_heading_custom(doc, "6.3 Sensitivity to neighbor cap k", 2)
    add_caption(
        doc,
        "Table 4. Minimum Tx for mean degree ≥8 across neighbor caps. For k=6 the degree-8 "
        "target is unreachable because stored degree cannot exceed k (entries marked —).",
    )
    add_table(
        doc,
        ["N", "k=6", "k=10 Min Tx%", "k=12 Min Tx%"],
        [
            ["50", "—", "20", "20"],
            ["100", "—", "15", "15"],
            ["150", "—", "10", "10"],
            ["200", "—", "10", "10"],
            ["300", "—", "8", "8"],
        ],
    )
    add_para(
        doc,
        "Thus k≈10 remains a practical operating point: large enough for stable maps, small "
        "enough to avoid distant weak neighbors. Increasing k from 10 to 12 does not reduce "
        "the energy-optimal Tx further in this matrix.",
    )

    add_heading_custom(doc, "6.4 Hard versus soft convergence", 2)
    add_para(
        doc,
        "A hard stop that requires an entire random batch to find zero violations can fail to "
        "trigger even when residuals are small, especially at high degree. Soft criteria "
        "(violation ≤1 px with adequate degree and connectivity) better match engineering "
        "stopping rules. Sparse, very low-Tx graphs may “hard-converge” trivially because few "
        "constraints exist; such cases are excluded from soft success by the degree≥6 rule.",
    )

    add_heading_custom(doc, "6.5 Prior directional-consistency observations (scoped)", 2)
    add_para(
        doc,
        "Earlier manuscript drafts reported relative-angle MAE below 6° on planned layouts "
        "across a range of network sizes, and directional correctness near 94.7% under an "
        "idealized k-NN model (≈89.5% under log-normal shadowing with σ≈4 dB). Those figures "
        "motivated the routing-centric thesis. In V02 we retain them only as scoped prior "
        "observations and do not mix them with the new density–Tx matrix without a single "
        "unified re-run. Baseline message-count comparisons from V01 are likewise treated as "
        "indicative, not re-certified here.",
        space_after=4,
    )
    add_para(
        doc,
        "Removed from V02: inconsistent iteration-count extrapolations (including ad-hoc "
        "linear-in-N rules and ~4000-iteration claims at N=1000) and extrapolations to "
        "10,000 nodes. Those statements were mutually inconsistent with the stated iteration "
        "cap and are not part of the auditable Results/ corpus.",
        space_after=10,
    )

    # ----- 7 Routing -----
    add_heading_custom(doc, "7 Integration with Direction-Aware Routing", 1)
    add_para(
        doc,
        "Given relative positions, a node computes the destination direction in the virtual "
        "frame and selects the neighbor minimizing angular deviation (greedy directional "
        "forwarding). Because similarity transforms preserve angles, an embedding that is "
        "consistent up to similarity supports the same angular decisions as a reference layout.",
    )
    add_para(
        doc,
        "Routing failures in greedy schemes are often topological voids (local minima), not "
        "pure localization error [1,14]. Perimeter/face recovery remains complementary. "
        "V02 does not claim new packet-level routing campaigns; integrating relative coordinates "
        "into GPSR-style stacks is future work. The density–Tx results matter for routing "
        "because connectivity and degree govern both embedding quality and forwarding options: "
        "energy-optimal Tx must still keep the graph usable for greedy progress.",
    )

    # ----- 8 Conclusion -----
    add_heading_custom(doc, "8 Conclusion, Limitations, and Future Work", 1)
    add_heading_custom(doc, "8.1 Conclusion", 2)
    add_para(
        doc,
        "Relative, reference-free embeddings aimed at directional consistency are a credible "
        "lightweight alternative to metric localization for direction-aware WSN routing. The "
        "asymmetric attraction rule offers clear local progress; global convergence remains "
        "empirically supported rather than fully proved. The new density–Tx study provides "
        "actionable deployment guidance: increase density to reduce Tx (and energy) while "
        "holding neighbor degree near 8–10 (k≈10).",
    )
    add_caption(
        doc,
        "Table 5. Key V02 findings (telecom perspective).",
    )
    add_table(
        doc,
        ["Finding", "Evidence", "Implication"],
        [
            [
                "Directional objective suffices in principle",
                "Scoped prior angular / directional results",
                "Avoid over-designing for metric accuracy",
            ],
            [
                "Min Tx falls with density",
                "Table 3 / Figs. 1–2 (600 trials)",
                "Deploy denser to save Tx energy",
            ],
            [
                "Degree target ≈8–10",
                "Fig. 3; k-sweep Table 4",
                "Cap neighbors; do not over-range",
            ],
            [
                "Soft stopping needed",
                "Section 6.4",
                "Use residual thresholds in practice",
            ],
            [
                "Normalization not yet decentralized",
                "Section 4.2",
                "Future gossip / local consensus",
            ],
        ],
    )

    add_heading_custom(doc, "8.2 Limitations", 2)
    limits = [
        "Global scale normalization in the current simulator.",
        "Density–Tx matrix uses planned-geometry constraints, fixed canvas, and N≤300.",
        "Hard finish flags under-report practical convergence; soft metrics are preferred.",
        "Packet-level routing and distributed normalization are not re-validated in V02.",
        "Assumes a working connected component; extremely sparse graphs remain out of scope.",
    ]
    for b in limits:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(b)
        set_run_font(run, size=11)

    add_heading_custom(doc, "8.3 Future work", 2)
    futures = [
        "Implement and measure gossip-based normalization [19] against the global baseline.",
        "Packet-level GPSR/perimeter experiments comparing relative vs. true coordinates on void rates, path stretch, and overhead.",
        "Map Tx% to dBm using the log-normal model (PL0, η, σ) for radio sizing worksheets.",
        "Extend density–Tx sweeps to larger N and irregular obstacle maps.",
    ]
    for b in futures:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(b)
        set_run_font(run, size=11)

    # ----- Declarations -----
    add_heading_custom(doc, "Declarations", 1)
    add_para(doc, "Competing interests: The authors declare no competing interests.", space_after=4)
    add_para(doc, "Funding: Not applicable.", space_after=4)
    add_para(
        doc,
        "Author contributions: Kanak Patel — conceptualization, methodology, software, experiments, "
        "writing. Himanshu S. Mazumdar — supervision of simulator design, density–Tx study guidance, "
        "manuscript revision.",
        space_after=4,
    )
    add_para(
        doc,
        "Data availability: Experiment CSVs, figures, and captions are provided under "
        "KanakPaperLocalization/Results/ (param_matrix.csv, results_raw.csv, "
        "results_summary_soft.csv, figures/, captions.md, paper_tables.md). Source prototypes: "
        "WsnLocalization2026-main and experiments/.",
        space_after=10,
    )

    # ----- References -----
    add_heading_custom(doc, "References", 1)
    refs = [
        "[1] Karp B, Kung HT (2000) GPSR: Greedy perimeter stateless routing for wireless networks. MobiCom, pp 243–254.",
        "[2] Patwari N et al. (2005) Locating the nodes: Cooperative localization in wireless sensor networks. IEEE Signal Process Mag 22(4):54–69.",
        "[3] Langendoen K, Reijers N (2003) Distributed localization in wireless sensor networks: A quantitative comparison. Comput Netw 43(4):499–518.",
        "[4] Shang Y, Ruml W, Zhang Y, Fromherz M (2003) Localization from mere connectivity. MobiHoc, pp 201–212.",
        "[5] Niculescu D, Nath B (2003) DV based positioning in ad hoc networks. Telecommun Syst 22(1–4):267–280.",
        "[6] He T et al. (2003) Range-free localization schemes for large scale sensor networks. MobiCom, pp 81–95.",
        "[7] Savvides A, Park H, Srivastava MB (2003) The n-hop multilateration primitive for node localization problems. MONET 8(4):443–451.",
        "[8] Doherty L, Pister KSJ, El Ghaoui L (2001) Convex position estimation in wireless sensor networks. INFOCOM.",
        "[9] Costa JA, Patwari N, Hero AO (2006) Distributed weighted-multidimensional scaling for node localization. ACM TOSN 2(1):39–64.",
        "[10] Biswas P et al. (2006) Semidefinite programming approaches for sensor network localization. IEEE TASE 3(4):360–371.",
        "[11] Zhang B et al. (2024) Accurate localization in LOS/NLOS channel coexistence scenarios based on heterogeneous knowledge graph inference. ACM TOSN 20(4):98.",
        "[12] Howard A, Mataric MJ, Sukhatme GS (2002) Localization for mobile robot teams using maximum likelihood estimation. IROS.",
        "[13] Hightower J, Borriello G (2001) Location systems for ubiquitous computing. IEEE Computer 34(8):57–66.",
        "[14] Kuhn F et al. (2003) Geometric ad-hoc routing: Of theory and practice. PODC, pp 63–72.",
        "[15] Basagni S et al. (1998) A distance routing effect algorithm for mobility (DREAM). MobiCom, pp 76–84.",
        "[16] Fonseca R et al. (2005) Beacon vector routing. NSDI.",
        "[17] Abraham I, Dolev D, Shavitt Y (2004) Routing with guaranteed delivery using virtual coordinates. INFOCOM.",
        "[18] Yu Y, Govindan R, Estrin D (2001) Geographical and energy-aware routing. UCLA-CSD TR-01-0023.",
        "[19] Boyd S et al. (2006) Randomized gossip algorithms. IEEE Trans Inf Theory 52(6):2508–2530.",
        "[20] Gupta P, Kumar PR (1998) Critical power for asymptotic connectivity in wireless networks. Stochastic Analysis, Control, Optimization and Applications.",
    ]
    for r in refs:
        add_para(doc, r, size=9, space_after=3)

    # Change log appendix
    add_heading_custom(doc, "Appendix A — Claim reconciliation checklist", 1)
    add_caption(
        doc,
        "Table 6. Claims removed, softened, or newly supported in V02.",
    )
    add_table(
        doc,
        ["V01 claim / issue", "V02 action"],
        [
            [
                "Fully distributed while using global normalization",
                "Marked Partial; limitation stated; gossip not asserted as done",
            ],
            [
                "Lyapunov / global convergence “guarantees”",
                "Lemma 1 local only; global labeled empirical",
            ],
            [
                "Inconsistent iteration-count extrapolations vs T_max",
                "Removed inconsistent scalability narrative",
            ],
            [
                "Extrapolation to 10,000 nodes",
                "Removed",
            ],
            [
                "Dummy gossip / packet-level routing results in response letter",
                "Not included as manuscript claims",
            ],
            [
                "Broken / unverifiable GitHub-only evidence",
                "Point to local Results/ + experiments/ artifacts",
            ],
            [
                "Density–Tx energy deployment assist",
                "Added Section 6.2–6.4 with figures/tables",
            ],
        ],
    )

    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    build()
